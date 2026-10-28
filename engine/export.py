"""A transcript as TXT, Markdown, SRT, VTT, DOCX or JSON."""

from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

FORMATS = ("txt", "md", "srt", "vtt", "docx", "json")
SOURCES = {"meet": "Meeting", "file": "File", "mic": "Microphone", "note": "Note"}
MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December")
SENTENCE = re.compile(r"(?<=[.!?…])\s+")


def clock(seconds: float, millis: str | None = None) -> str:
    """00:12, 1:02:03; with millis="," or "." the subtitle form 00:00:12,345."""
    total = round(seconds * 1000)  # rounded once, so 0.9996 s carries into 00:00:01,000
    whole, ms = divmod(total, 1000) if millis else (int(seconds), 0)
    h, m, s = whole // 3600, whole % 3600 // 60, whole % 60
    if millis:
        return f"{h:02}:{m:02}:{s:02}{millis}{ms:03}"
    return f"{h}:{m:02}:{s:02}" if h else f"{m:02}:{s:02}"


def length(seconds: float) -> str:
    minutes = round(seconds / 60)
    if minutes < 1:
        return f"{int(seconds)} s"
    return f"{minutes // 60} h {minutes % 60} min" if minutes >= 60 else f"{minutes} min"


def when(created: str) -> str:
    d = datetime.fromisoformat(created)
    return f"{d.day} {MONTHS[d.month - 1]} {d.year}, {d:%H:%M}"


def name(t: dict, key: str) -> str:
    return t.get("speakers", {}).get(key, {}).get("name", key)


def participants(t: dict) -> list[str]:
    seen = []
    for u in t.get("utterances", []):
        if u["speaker"] not in seen:
            seen.append(u["speaker"])
    return [name(t, k) for k in seen]


def meta_line(t: dict) -> str:
    return " · ".join([when(t["created"]), length(t.get("duration", 0)), SOURCES.get(t.get("source", ""), "")])


def to_txt(t: dict) -> str:
    lines = [t.get("title", ""), meta_line(t), ""]
    lines += [f"[{clock(u['start'])}] {name(t, u['speaker'])}: {u['text']}" for u in t.get("utterances", [])]
    return "\n".join(lines) + "\n"


def to_md(t: dict, heading: bool = True) -> str:
    out = [f"# {t.get('title', '')}", ""] if heading else []
    out += [f"*{meta_line(t)}*", ""]
    people = participants(t)
    if len(people) > 1:
        out += [f"**Speakers:** {', '.join(people)}", ""]
    previous = None
    for u in t.get("utterances", []):
        if len(people) > 1 and u["speaker"] != previous:
            out.append(f"**{name(t, u['speaker'])}** `{clock(u['start'])}`  ")
        elif len(people) <= 1:
            out.append(f"`{clock(u['start'])}`  ")
        out += [u["text"], ""]
        previous = u["speaker"]
    return "\n".join(out)


def cues(t: dict) -> list[tuple[float, float, str]]:
    """Utterances cut into subtitles of a sentence or so, the time shared out by length."""
    out = []
    for u in t.get("utterances", []):
        parts = [p for p in SENTENCE.split(u["text"]) if p]
        span = max(u["end"] - u["start"], 0.1)
        chars = sum(len(p) for p in parts) or 1
        start = u["start"]
        speaker = name(t, u["speaker"])
        multi = len(t.get("speakers", {})) > 1
        for p in parts:
            end = start + span * len(p) / chars
            out.append((start, end, f"{speaker}: {p}" if multi else p))
            start = end
    return out


def to_srt(t: dict) -> str:
    return "\n".join(
        f"{i}\n{clock(a, ',')} --> {clock(b, ',')}\n{text}\n" for i, (a, b, text) in enumerate(cues(t), 1)
    )


def to_vtt(t: dict) -> str:
    return "WEBVTT\n\n" + "\n".join(f"{clock(a, '.')} --> {clock(b, '.')}\n{text}\n" for a, b, text in cues(t))


def to_docx(t: dict, path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.add_heading(t.get("title", ""), level=1)
    meta = doc.add_paragraph(meta_line(t))
    meta.runs[0].font.color.rgb = RGBColor(0x6E, 0x76, 0x81)
    people = participants(t)
    if len(people) > 1:
        doc.add_paragraph(f"Speakers: {', '.join(people)}")
    for u in t.get("utterances", []):
        p = doc.add_paragraph()
        stamp = p.add_run(f"{clock(u['start'])}  ")
        stamp.font.color.rgb = RGBColor(0x6E, 0x76, 0x81)
        stamp.font.size = Pt(9)
        if len(people) > 1:
            p.add_run(f"{name(t, u['speaker'])}: ").bold = True
        p.add_run(u["text"])
    doc.save(str(path))


def write(t: dict, fmt: str, path: Path) -> None:
    if fmt == "docx":
        to_docx(t, path)
        return
    text = {
        "txt": to_txt,
        "md": to_md,
        "srt": to_srt,
        "vtt": to_vtt,
        "json": lambda t: json.dumps(t, ensure_ascii=False, indent=1),
    }[fmt](t)
    path.write_text(text, "utf-8")
